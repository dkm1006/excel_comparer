"""Core comparison logic.

The :class:`ExcelComparer` loads a *golden* workbook once and can then be
asked to compare any number of other workbooks against it.  Comparison is
strict: sheets are matched by exact name, cells by coordinate (``A1`` ↔
``A1``).  Cached computed values are used (``data_only=True``); formulas
themselves are not inspected.
"""

import math
import itertools
from pathlib import Path
from typing import Any, Iterable

import openpyxl
from openpyxl import load_workbook
from openpyxl.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from excel_comparer.models import DiffCategory, Difference


# Attributes that we compare cell-by-cell when at least one side is non-empty.
# Each entry is (dotted attribute path on ``Cell``, human-readable label).
_FORMAT_ATTRIBUTES: tuple[tuple[str, str], ...] = (
    ("font.name", "font.name"),
    ("font.size", "font.size"),
    ("font.bold", "font.bold"),
    ("font.italic", "font.italic"),
    ("font.underline", "font.underline"),
    ("font.color.rgb", "font.color.rgb"),
    ("fill.patternType", "fill.patternType"),
    ("fill.fgColor.rgb", "fill.fgColor.rgb"),
    ("fill.bgColor.rgb", "fill.bgColor.rgb"),
    ("number_format", "number_format"),
    ("alignment.horizontal", "alignment.horizontal"),
    ("alignment.vertical", "alignment.vertical"),
    ("alignment.wrap_text", "alignment.wrap_text"),
    ("border.left.style", "border.left.style"),
    ("border.left.color.rgb", "border.left.color.rgb"),
    ("border.right.style", "border.right.style"),
    ("border.right.color.rgb", "border.right.color.rgb"),
    ("border.top.style", "border.top.style"),
    ("border.top.color.rgb", "border.top.color.rgb"),
    ("border.bottom.style", "border.bottom.style"),
    ("border.bottom.color.rgb", "border.bottom.color.rgb"),
)


class ExcelComparer:
    """Compare other Excel workbooks against a fixed *golden* workbook.

    Parameters
    ----------
    golden_path:
        Path to the golden ``.xlsx`` file.  Loaded eagerly.
    float_tolerance:
        Absolute tolerance applied when comparing floating-point cell values.
        Defaults to ``0.0`` (exact equality).
    """
    DEFAULT_DIFF_CATEGORIES = (DiffCategory.VALUE, DiffCategory.SHEET_MISSING, DiffCategory.SHEET_EXTRA)

    def __init__(
        self,
        golden_path: str | Path,
        float_tolerance: float = 0.0,
        diff_categories: tuple[DiffCategory] = DEFAULT_DIFF_CATEGORIES
    ) -> None:
        self.golden_path: Path = Path(golden_path)
        self.float_tolerance: float = float_tolerance
        self.diff_categories: tuple[DiffCategory] = diff_categories
        self.golden: Workbook = load_workbook(
            filename=str(self.golden_path), data_only=True
        )

    def compare(self, *others: str | Path) -> dict[str, list[Difference]]:
        """Compare one or more workbooks against the golden workbook.

        Returns a dict ``{path_str: [Difference, ...]}`` with one entry
        per argument, in the order they were given.
        """
        results: dict[str, list[Difference]] = {}
        for other in others:
            other_path = Path(other)
            results[str(other)] = self._compare_one(other_path)
        return results

    def _compare_one(self, other_path: Path) -> list[Difference]:
        other_wb: Workbook = load_workbook(filename=str(other_path), data_only=True)

        diffs: list[Difference] = []

        # Sheet-level diffs
        if DiffCategory.SHEET_MISSING in self.diff_categories or DiffCategory.SHEET_EXTRA in self.diff_categories:
            diffs.extend(self._compare_sheetsets(other_wb.sheetnames))

        # Per-sheet diffs for sheets present on both sides.
        for name in self.golden.sheetnames:
            if name in other_wb.sheetnames:
                diffs.extend(self._compare_sheet(self.golden[name], other_wb[name]))
        return diffs

    def _compare_sheetsets(
        self, other_sheets: Iterable[str]
    ) -> Iterable[Difference]:
        """Compare the sets of sheet names in the golden and other workbooks."""
        golden_set = set(self.golden.sheetnames)
        other_set = set(other_sheets)
        for name in sorted(golden_set - other_set):
            yield Difference(category=DiffCategory.SHEET_MISSING, sheet=name)
        for name in sorted(other_set - golden_set):
            yield Difference(category=DiffCategory.SHEET_EXTRA, sheet=name)

    def _compare_sheet(
        self, g_ws: Worksheet, o_ws: Worksheet
    ) -> Iterable[Difference]:
        sheet_name = g_ws.title
        # TODO: Check if max_row and max_column can be None. If yes do g_ws.max_row or 0
        max_row = max(g_ws.max_row, o_ws.max_row)
        max_col = max(g_ws.max_column, o_ws.max_column)

        diffs: list[Difference] = []

        # TODO: Check if index starts at 1 for both row and column in openpyxl. 
        # If yes, the current loop is correct. 
        # If it starts at 0, adjust the range to start from 0 and go to max_row and max_col.
        row_col_combinations = itertools.product(
            range(1, max_row + 1), range(1, max_col + 1)
        )
        for row, col in row_col_combinations:
            g_cell = g_ws.cell(row=row, column=col)
            o_cell = o_ws.cell(row=row, column=col)

            if DiffCategory.VALUE in self.diff_categories:
                if diff := self._compare_values(sheet_name, g_cell, o_cell):
                    diffs.append(diff)

            # Format diffs only when at least one side actually has a value.
            if DiffCategory.FORMAT in self.diff_categories and not (_is_empty(g_cell) and _is_empty(o_cell)):
                diffs.extend(self._compare_formats(sheet_name, g_cell, o_cell))

        # Merged-cell ranges
        if DiffCategory.MERGED_CELL in self.diff_categories:
            diffs.extend(self._compare_merged_cells(g_ws, o_ws))
    
        return diffs

    def _compare_formats(self, sheet_name: str, g_cell: openpyxl.cell.Cell, o_cell: openpyxl.cell.Cell):
        for attr_path, label in _FORMAT_ATTRIBUTES:
            g_attr = _get_attr_path(g_cell, attr_path)
            o_attr = _get_attr_path(o_cell, attr_path)
            if g_attr != o_attr:
                yield Difference(
                        category=DiffCategory.FORMAT,
                        sheet=sheet_name,
                        cell=g_cell.coordinate,
                        attribute=label,
                        golden=g_attr,
                        other=o_attr
                    )

    def _compare_values(self, sheet_name: str, g_cell: openpyxl.cell.Cell, o_cell: openpyxl.cell.Cell) -> Difference|None:
        """Compare two cell values, honouring an optional float tolerance."""
        diff = None
        if not _values_equal(g_cell, o_cell, self.float_tolerance):
            diff = Difference(
                category=DiffCategory.VALUE,
                sheet=sheet_name,
                cell=g_cell.coordinate,
                golden=g_cell.value,
                other=o_cell.value
            )
        return diff

    def _compare_merged_cells(self, g_ws: Worksheet, o_ws: Worksheet) -> Iterable[Difference]:
        # TODO: Check this again
        g_ranges = {str(r) for r in g_ws.merged_cells.ranges}
        o_ranges = {str(r) for r in o_ws.merged_cells.ranges}
        for r in sorted(g_ranges - o_ranges):
            yield Difference(
                    category=DiffCategory.MERGED_CELL,
                    sheet=g_ws.title,
                    attribute="merged_range",
                    golden=r,
                    other=None
                )
        for r in sorted(o_ranges - g_ranges):
            yield Difference(
                    category=DiffCategory.MERGED_CELL,
                    sheet=g_ws.title,
                    attribute="merged_range",
                    golden=None,
                    other=r
                )


def _is_empty(cell: openpyxl.cell.Cell) -> bool:
    """Treat ``None`` and the empty string as "no value"."""
    return cell.value is None or cell.value == ""


def _values_equal(a: openpyxl.cell.Cell, b: openpyxl.cell.Cell, float_tolerance: float) -> bool:
    """Compare two cell values, honouring an optional float tolerance."""
    if _is_empty(a) and _is_empty(b):
        is_equal = True
    elif _is_empty(a) or _is_empty(b):
        is_equal = False
    elif isinstance(a.value, float) or isinstance(b.value, float):
        try:
            is_equal = math.isclose(float(a.value), float(b.value), abs_tol=float_tolerance, rel_tol=0.0)
        except (TypeError, ValueError):
            is_equal = False
    else:
        is_equal = a.value == b.value
    return is_equal


def _get_attr_path(obj: Any, path: str) -> Any:
    """Resolve a dotted attribute path, returning ``None`` if any step is missing."""
    current = obj
    for part in path.split("."):
        if current is None:
            return None
        current = getattr(current, part, None)
    return current
