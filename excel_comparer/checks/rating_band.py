"""Rating-band check.

Verifies that every non-empty cell inside the configured ``ranges`` holds
a numeric value within ``[min_value, max_value]`` (optionally restricted
to whole integers). Empty cells are *not* reported here — that's the
:class:`RequiredCellsCheck`'s job.

The check runs on every sheet listed in ``ranges``. If ``ranges`` is
empty or missing the check is a no-op.

Configuration (TOML)
--------------------

.. code-block:: toml

    [[checks]]
    name = "rating_band"
    min_value = 1
    max_value = 4
    integers_only = true

    [checks.ranges]
    Impacts = ["O3:O", "P3:P", "Q3:Q", "S3:S"]
"""

from typing import ClassVar

from openpyxl.utils import get_column_letter

from excel_comparer.checks.base import CheckContext, register_check
from excel_comparer.checks.helpers import (
    find_last_data_row,
    is_empty,
    is_in_range,
    normalize_column,
)
from excel_comparer.config import SheetRanges
from excel_comparer.models import DiffCategory, Difference


@register_check
class RatingBandCheck:
    """Assert that every cell in ``ranges`` holds a number in ``[min_value, max_value]``."""

    name: ClassVar[str] = "rating_band"
    description: ClassVar[str] = (
        "Asserts that the values in every cell of the configured ranges "
        "are within a numeric range. Optionally restricts to integers."
    )

    def __init__(
        self,
        *,
        min_value: float = 1,
        max_value: float = 4,
        integers_only: bool = True,
        data_start_row: int = 3,
        row_anchor_column: int | str | None = None,
        ranges: dict[str, SheetRanges] | None = None,
    ) -> None:
        self.min_value = min_value
        self.max_value = max_value
        self.integers_only = bool(integers_only)
        self.data_start_row = int(data_start_row)
        self.row_anchor_column: int | None = (
            normalize_column(row_anchor_column) if row_anchor_column else None
        )
        self.ranges: dict[str, SheetRanges] = ranges or {}

    def run(self, ctx: CheckContext) -> list[Difference]:
        diffs: list[Difference] = []
        for sheet_name, sr in self.ranges.items():
            if sr.skip or sheet_name not in ctx.workbook.sheetnames:
                continue
            ws = ctx.workbook[sheet_name]

            # Derive the scan column span from the configured ranges so we
            # don't scan the whole sheet looking for the last data row.
            scan_min_col, scan_max_col = sr.clamp_cols(1, ws.max_column)
            if scan_min_col > scan_max_col:
                continue
            last_row = find_last_data_row(
                ws,
                start_row=self.data_start_row,
                scan_min_col=scan_min_col,
                scan_max_col=scan_max_col,
            )
            if last_row < self.data_start_row:
                continue

            anchor_empty: dict[int, bool] = {}

            for row, col in sr.iter_cells(
                row_floor=self.data_start_row,
                row_ceiling=last_row,
                col_ceiling=ws.max_column,
            ):
                if self.row_anchor_column is not None:
                    if row not in anchor_empty:
                        anchor_empty[row] = is_empty(
                            ws.cell(row=row, column=self.row_anchor_column).value
                        )
                    if anchor_empty[row]:
                        continue
                value = ws.cell(row=row, column=col).value
                if is_empty(value):
                    continue
                if is_in_range(
                    value,
                    self.min_value,
                    self.max_value,
                    integers_only=self.integers_only,
                ):
                    continue
                coord = f"{get_column_letter(col)}{row}"
                expected = (
                    f"integer in [{self.min_value}, {self.max_value}]"
                    if self.integers_only
                    else f"value in [{self.min_value}, {self.max_value}]"
                )
                diffs.append(
                    Difference(
                        category=DiffCategory.RATING_OUT_OF_RANGE,
                        sheet=sheet_name,
                        cell=coord,
                        golden=expected,
                        other=value,
                        check_name=self.name,
                    )
                )
        return diffs
