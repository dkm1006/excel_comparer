"""Baseline (golden-workbook) comparison check.

Compares a candidate workbook cell-by-cell against a fixed baseline file
and emits :class:`Difference` entries for value, format, merged-cell and
sheet-set discrepancies. This is the original behaviour of the package,
re-cast as a :class:`Check` so it sits alongside other (non-baseline)
checks in the configuration.

Matching is strict: sheets by exact name, cells by coordinate (``A1`` ↔
``A1``). Cached computed values are used (``data_only=True``); formulas
themselves are not inspected.

Configuration (TOML)
--------------------

.. code-block:: toml

    [[checks]]
    name = "baseline_comparison"
    float_tolerance = 0.0
    categories = ["VALUE", "FORMAT", "MERGED_CELL", "SHEET_MISSING", "SHEET_EXTRA"]

    # Optional: restrict (or skip) per-sheet comparison. Sheets *not*
    # listed are compared in full.
    [checks.ranges]
    Impacts = "A1:Z"
    "Lookup tables" = { skip = true }
"""

import math
from pathlib import Path
from typing import Any, ClassVar, Iterable

import openpyxl
from openpyxl import load_workbook
from openpyxl.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from excel_comparer.checks.base import CheckContext, register_check
from excel_comparer.config import SheetRange, SheetRanges
from excel_comparer.models import DiffCategory, Difference


# Attributes that we compare cell-by-cell when at least one side is non-empty.
# Each entry is (dotted attribute path on ``Cell``, human-readable label).
FORMAT_ATTRIBUTES: tuple[tuple[str, str], ...] = (
    ("font.name", "font.name"),
    ("font.size", "font.size"),
    ("font.bold", "font.bold"),
    ("font.italic", "font.italic"),
    ("font.underline", "font.underline"),
    ("font.color.rgb", "font.color.rgb"),
    ("fill.patternType", "fill.patternType"),
    ("fill.fgColor.rgb", "fill.fgColor.rgb"),
    ("fill.bgColor.rgb", "fill.bgColor.rgb"),
    ("number_format", "number_format"),
    ("alignment.horizontal", "alignment.horizontal"),
    ("alignment.vertical", "alignment.vertical"),
    ("alignment.wrap_text", "alignment.wrap_text"),
    ("border.left.style", "border.left.style"),
    ("border.left.color.rgb", "border.left.color.rgb"),
    ("border.right.style", "border.right.style"),
    ("border.right.color.rgb", "border.right.color.rgb"),
    ("border.top.style", "border.top.style"),
    ("border.top.color.rgb", "border.top.color.rgb"),
    ("border.bottom.style", "border.bottom.style"),
    ("border.bottom.color.rgb", "border.bottom.color.rgb"),
)


DEFAULT_CATEGORIES: tuple[DiffCategory, ...] = (
    DiffCategory.VALUE,
    DiffCategory.SHEET_MISSING,
    DiffCategory.SHEET_EXTRA,
)

BASELINE_CATEGORIES: frozenset[DiffCategory] = frozenset({
    DiffCategory.VALUE,
    DiffCategory.FORMAT,
    DiffCategory.MERGED_CELL,
    DiffCategory.SHEET_MISSING,
    DiffCategory.SHEET_EXTRA,
})


@register_check
class BaselineComparisonCheck:
    """Compare candidate workbooks against a fixed *baseline* workbook.

    Parameters
    ----------
    baseline_path:
        Path to the baseline ``.xlsx`` file. Loaded lazily on first
        ``run()`` (and cached) so the check can be constructed before the
        CLI knows the path. May be ``None`` only if the orchestrator will
        inject it later by setting :attr:`baseline_path` directly.
    float_tolerance:
        Absolute tolerance applied when comparing floating-point cell
        values. Defaults to ``0.0`` (exact equality).
    categories:
        Categories to report (any subset of ``VALUE``, ``FORMAT``,
        ``MERGED_CELL``, ``SHEET_MISSING``, ``SHEET_EXTRA``).
    ranges:
        Optional mapping ``sheet_name -> SheetRanges`` restricting the
        compared region for that sheet (or marking it ``skip=True``).
        Sheets not present in the mapping are compared in full. The
        config layer coerces TOML strings / lists / dicts into
        :class:`SheetRanges` instances before they reach this check.
    """

    name: ClassVar[str] = "baseline_comparison"
    description: ClassVar[str] = (
        "Compare each candidate workbook cell-by-cell against a baseline "
        "reference workbook and report value, format, merged-cell and "
        "sheet-set differences."
    )

    def __init__(
        self,
        baseline_path: str | Path | None = None,
        *,
        float_tolerance: float = 0.0,
        categories: Iterable[str | DiffCategory] = DEFAULT_CATEGORIES,
        ranges: dict[str, SheetRanges] | None = None,
        # Accepted for symmetry with other checks (top-level [defaults]
        # injection); unused by this check.
        header_row: int | None = None,  # noqa: ARG002
        row_anchor_column: int | str | None = None,  # noqa: ARG002
    ) -> None:
        self.baseline_path: Path | None = Path(baseline_path) if baseline_path else None
        self.float_tolerance: float = float(float_tolerance)
        self.categories: tuple[DiffCategory, ...] = _coerce_categories(categories)
        self.ranges: dict[str, SheetRanges] = ranges or {}
        self._baseline_wb: Workbook | None = None

    def run(self, ctx: CheckContext) -> list[Difference]:
        baseline = self._load_baseline()
        candidate = ctx.workbook

        diffs: list[Difference] = []
        diffs.extend(self._compare_sheetsets(baseline, candidate.sheetnames))

        for sheet_name in baseline.sheetnames:
            if sheet_name in candidate.sheetnames:
                sheet_ranges = self.ranges.get(sheet_name)
                diffs.extend(
                    self._compare_sheet(
                        baseline[sheet_name], candidate[sheet_name], sheet_ranges
                    )
                )
        return diffs

    # ------------------------------------------------------------------ #
    # Internals                                                          #
    # ------------------------------------------------------------------ #

    def _load_baseline(self) -> Workbook:
        if self.baseline_path is None:
            raise ValueError(
                f"{type(self).__name__}: baseline_path is not set. Provide it "
                "via config or the CLI's positional `baseline` argument."
            )
        if self._baseline_wb is None:
            self._baseline_wb = load_workbook(
                filename=str(self.baseline_path), data_only=True
            )
        return self._baseline_wb

    def _compare_sheetsets(
        self, baseline: Workbook, candidate_sheets: Iterable[str]
    ) -> Iterable[Difference]:
        baseline_set = set(baseline.sheetnames)
        candidate_set = set(candidate_sheets)
        if DiffCategory.SHEET_MISSING in self.categories:
            for name in sorted(baseline_set - candidate_set):
                yield Difference(category=DiffCategory.SHEET_MISSING, sheet=name, check_name=self.name)
        if DiffCategory.SHEET_EXTRA in self.categories:
            for name in sorted(candidate_set - baseline_set):
                yield Difference(category=DiffCategory.SHEET_EXTRA, sheet=name, check_name=self.name)

    def _compare_sheet(
        self,
        g_ws: Worksheet,
        o_ws: Worksheet,
        sheet_ranges: SheetRanges | None,
    ) -> Iterable[Difference]:
        sheet_name = g_ws.title

        # Skip the whole sheet if the config said so.
        if sheet_ranges is not None and sheet_ranges.skip:
            return []

        diffs: list[Difference] = []
        max_row = max(g_ws.max_row, o_ws.max_row)
        max_col = max(g_ws.max_column, o_ws.max_column)

        has_cell_bound_category = any(
            cat in _CELL_BOUND_BASELINE_CATEGORIES for cat in self.categories
        )
        if has_cell_bound_category and max_row >= 1 and max_col >= 1:
            # Sheets not listed in ranges -> compare in full.
            sr = sheet_ranges if (sheet_ranges and sheet_ranges.ranges) else _FULL_SHEET
            for row, col in sr.iter_cells(
                row_floor=1,
                row_ceiling=max_row,
                col_floor=1,
                col_ceiling=max_col,
            ):
                g_cell = g_ws.cell(row=row, column=col)
                o_cell = o_ws.cell(row=row, column=col)

                if DiffCategory.VALUE in self.categories:
                    if diff := self._compare_values(sheet_name, g_cell, o_cell):
                        diffs.append(diff)

                # Format diffs only when at least one side actually has a value.
                if (
                    DiffCategory.FORMAT in self.categories
                    and not (_cell_is_empty(g_cell) and _cell_is_empty(o_cell))
                ):
                    diffs.extend(self._compare_formats(sheet_name, g_cell, o_cell))

        # Merged-cell ranges (whole sheet; not row/col limited).
        if DiffCategory.MERGED_CELL in self.categories:
            diffs.extend(self._compare_merged_cells(g_ws, o_ws))

        return diffs

    def _compare_formats(
        self, sheet_name: str, g_cell: openpyxl.cell.Cell, o_cell: openpyxl.cell.Cell
    ) -> Iterable[Difference]:
        for attr_path, label in FORMAT_ATTRIBUTES:
            g_attr = _get_attr_path(g_cell, attr_path)
            o_attr = _get_attr_path(o_cell, attr_path)
            if g_attr != o_attr:
                yield Difference(
                    category=DiffCategory.FORMAT,
                    sheet=sheet_name,
                    cell=g_cell.coordinate,
                    attribute=label,
                    golden=g_attr,
                    other=o_attr,
                    check_name=self.name
                )

    def _compare_values(
        self, sheet_name: str, g_cell: openpyxl.cell.Cell, o_cell: openpyxl.cell.Cell
    ) -> Difference | None:
        if _values_equal(g_cell, o_cell, self.float_tolerance):
            return None
        return Difference(
            category=DiffCategory.VALUE,
            sheet=sheet_name,
            cell=g_cell.coordinate,
            golden=g_cell.value,
            other=o_cell.value,
            check_name=self.name
        )

    def _compare_merged_cells(
        self, g_ws: Worksheet, o_ws: Worksheet
    ) -> Iterable[Difference]:
        g_ranges = {str(r) for r in g_ws.merged_cells.ranges}
        o_ranges = {str(r) for r in o_ws.merged_cells.ranges}
        for r in sorted(g_ranges - o_ranges):
            yield Difference(
                category=DiffCategory.MERGED_CELL,
                sheet=g_ws.title,
                attribute=r,
                golden=r,
                other=None,
                check_name=self.name
            )
        for r in sorted(o_ranges - g_ranges):
            yield Difference(
                category=DiffCategory.MERGED_CELL,
                sheet=g_ws.title,
                attribute=r,
                golden=None,
                other=r,
                check_name=self.name
            )


# --------------------------------------------------------------------------- #
# Module-level helpers                                                        #
# --------------------------------------------------------------------------- #


_CELL_BOUND_BASELINE_CATEGORIES = frozenset({DiffCategory.VALUE, DiffCategory.FORMAT})

# Sentinel SheetRanges used for sheets that aren't listed in the config:
# a single fully-unbounded sub-range so iter_cells walks the whole sheet
# (clamped to the union of both workbooks' used ranges by the caller).
_FULL_SHEET = SheetRanges(ranges=(SheetRange(),))


def _coerce_categories(raw: Iterable[str | DiffCategory]) -> tuple[DiffCategory, ...]:
    return tuple(DiffCategory(str(item)) for item in raw)


def _cell_is_empty(cell: openpyxl.cell.Cell) -> bool:
    """Treat ``None`` and the empty string as "no value"."""
    return cell.value is None or cell.value == ""


def _values_equal(
    a: openpyxl.cell.Cell, b: openpyxl.cell.Cell, float_tolerance: float
) -> bool:
    """Compare two cell values, honouring an optional float tolerance."""
    if _cell_is_empty(a) and _cell_is_empty(b):
        return True
    if _cell_is_empty(a) or _cell_is_empty(b):
        return False
    if isinstance(a.value, float) or isinstance(b.value, float):
        try:
            return math.isclose(
                float(a.value),
                float(b.value),
                abs_tol=float_tolerance,
                rel_tol=0.0,
            )
        except (TypeError, ValueError):
            return False
    return a.value == b.value


def _get_attr_path(obj: Any, path: str) -> Any:
    """Resolve a dotted attribute path, returning ``None`` if any step is missing."""
    current = obj
    for part in path.split("."):
        if current is None:
            return None
        current = getattr(current, part, None)
    return current
