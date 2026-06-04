"""Selection-required check.

For each data row, asserts that within each configured *column group*,
not every cell holds the configured ``forbidden_value`` (case-insensitive
after ``str().strip()``). Used in the IRO survey to enforce that at
least one value-chain position and at least one time horizon are
selected (the form uses ``"Yes"`` / ``"No"`` strings; an all-``"No"``
row is treated as a missing selection).

The resulting :class:`Difference` is anchored on the first cell of the
column group so the exporter highlights and the doc column line up with
the offending row.

The check runs on every sheet listed in ``ranges``. ``ranges`` restricts
which **rows** are visited; column groups define their own column spans.
A column group whose anchor column is excluded by the sheet's range is
skipped. If ``ranges`` is empty or missing the check is a no-op.

Configuration (TOML)
--------------------

.. code-block:: toml

    [[checks]]
    name = "selection_required"
    forbidden_value = "No"
    header_row = 2
    column_groups = [
        { label = "value_chain",  min_col = "I", max_col = "K" },
        { label = "time_horizon", min_col = "L", max_col = "N" },
    ]

    [checks.ranges]
    Impacts = "A3:Z"
"""

from dataclasses import dataclass
from typing import Any, ClassVar, Iterable

from openpyxl.utils import get_column_letter

from excel_comparer.checks.base import CheckContext, register_check
from excel_comparer.checks.helpers import (
    all_equal_case_insensitive,
    find_last_data_row,
    is_empty,
    normalize_column,
)
from excel_comparer.config import SheetRanges
from excel_comparer.models import DiffCategory, Difference


@dataclass(frozen=True)
class ColumnGroup:
    """One range of columns that must contain at least one non-forbidden value."""

    label: str
    min_col: int     # 1-based
    max_col: int     # 1-based, inclusive


@register_check
class SelectionRequiredCheck:
    """Assert each row picks at least one value across configured column groups."""

    name: ClassVar[str] = "selection_required"
    description: ClassVar[str] = (
        "Validates that within each configured column group, at least one cell "
        "in every data row holds a value other than the configured forbidden "
        "value (e.g. that at least one value-chain position is chosen)."
    )

    def __init__(
        self,
        *,
        column_groups: Iterable[dict[str, Any] | ColumnGroup],
        forbidden_value: str = "No",
        header_row: int = 2,
        row_anchor_column: int | str | None = None,
        scan_min_col: int | str = "B",
        scan_max_col: int | str = "G",
        ranges: dict[str, SheetRanges] | None = None,
    ) -> None:
        self.column_groups: list[ColumnGroup] = [
            _coerce_group(g) for g in column_groups
        ]
        self.forbidden_value = str(forbidden_value)
        self.header_row = header_row
        self.row_anchor_column: int | None = (
            normalize_column(row_anchor_column) if row_anchor_column else None
        )
        self.scan_min_col = normalize_column(scan_min_col)
        self.scan_max_col = normalize_column(scan_max_col)
        self.ranges: dict[str, SheetRanges] = ranges or {}

    def run(self, ctx: CheckContext) -> list[Difference]:
        diffs: list[Difference] = []
        for sheet_name, sr in self.ranges.items():
            if sr.skip or sheet_name not in ctx.workbook.sheetnames:
                continue
            ws = ctx.workbook[sheet_name]
            last_row = find_last_data_row(
                ws,
                start_row=self.header_row + 1,
                scan_min_col=self.scan_min_col,
                scan_max_col=self.scan_max_col,
            )
            lo, hi = sr.clamp_rows(self.header_row + 1, last_row)
            # Skip any group whose anchor (min_col) is outside the configured columns.
            active_groups = [g for g in self.column_groups if sr.col_in_range(g.min_col)]
            for row in range(lo, hi + 1):
                if not sr.row_in_range(row):
                    continue
                if self.row_anchor_column is not None and is_empty(
                    ws.cell(row=row, column=self.row_anchor_column).value
                ):
                    continue
                for group in active_groups:
                    values = [
                        ws.cell(row=row, column=col).value
                        for col in range(group.min_col, group.max_col + 1)
                    ]
                    if all_equal_case_insensitive(values, self.forbidden_value):
                        anchor = f"{get_column_letter(group.min_col)}{row}"
                        span = (
                            f"{get_column_letter(group.min_col)}-"
                            f"{get_column_letter(group.max_col)}"
                        )
                        diffs.append(
                            Difference(
                                category=DiffCategory.NO_SELECTION,
                                sheet=sheet_name,
                                cell=anchor,
                                attribute=group.label,
                                golden=f"at least one value other than {self.forbidden_value!r}",
                                other=f"all {self.forbidden_value!r} across columns {span}",
                                check_name=self.name,
                            )
                        )
        return diffs


def _coerce_group(spec: dict[str, Any] | ColumnGroup) -> ColumnGroup:
    if isinstance(spec, ColumnGroup):
        return spec
    if not isinstance(spec, dict):
        raise ValueError(
            f"selection_required.column_groups entry must be a table, got {type(spec).__name__}"
        )
    min_col = normalize_column(spec["min_col"])
    max_col = normalize_column(spec["max_col"])
    if max_col < min_col:
        raise ValueError(
            f"selection_required column group has max_col < min_col: {spec!r}"
        )
    label = str(spec.get("label") or f"{get_column_letter(min_col)}-{get_column_letter(max_col)}")
    return ColumnGroup(label=label, min_col=min_col, max_col=max_col)
