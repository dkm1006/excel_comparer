"""Irremediable-character check.

Asserts that the score in ``score_column`` is consistent with the
categorical type in ``type_column`` for every data row:

* When the type matches ``positive_value`` (default ``"Positive Impact"``),
  the score must be the *auto-correction value* (default ``"-"``).
  If it isn't, the check emits a diff whose ``correction`` field carries
  the expected value — the exporter applies the correction when it
  writes the annotated workbook (this is the only check that auto-corrects).
* When the type matches ``negative_value`` (default ``"Negative Impact"``),
  the score must be an integer in ``valid_scores``. Otherwise an
  ``INCONSISTENT_TYPE_SCORE`` diff is emitted.

Rows whose type cell is empty are skipped (no signal to validate against).

The check runs on every sheet listed in ``ranges``. ``ranges`` gates the
row range; the type/score columns must also lie inside the configured
columns for a row to be inspected. If ``ranges`` is empty or missing the
check is a no-op.

Configuration (TOML)
--------------------

.. code-block:: toml

    [[checks]]
    name = "irremediable_character"
    type_column = "G"
    score_column = "Q"
    positive_value = "Positive Impact"
    negative_value = "Negative Impact"
    auto_correction_value = "-"
    valid_scores = [1, 2, 3, 4]
    header_row = 2

    [checks.ranges]
    Impacts = "A3:Z"
"""

from typing import ClassVar, Iterable

from openpyxl.utils import get_column_letter

from excel_comparer.checks.base import CheckContext, register_check
from excel_comparer.checks.helpers import (
    find_last_data_row,
    is_empty,
    normalize_column,
)
from excel_comparer.config import SheetRanges
from excel_comparer.models import DiffCategory, Difference


@register_check
class IrremediableCharacterCheck:
    """Validate that a categorical type column drives the contents of a score column."""

    name: ClassVar[str] = "irremediable_character"
    description: ClassVar[str] = (
        "Asserts that the selected score for 'Irremediable Character' aligns "
        "with the impact type. For 'Positive Impact', the score is "
        "auto-corrected to '-'."
    )

    def __init__(
        self,
        *,
        type_column: int | str,
        score_column: int | str,
        positive_value: str = "Positive Impact",
        negative_value: str = "Negative Impact",
        auto_correction_value: str = "-",
        valid_scores: Iterable[int] = (1, 2, 3, 4),
        header_row: int = 2,
        row_anchor_column: int | str | None = None,
        scan_min_col: int | str = "B",
        scan_max_col: int | str = "G",
        ranges: dict[str, SheetRanges] | None = None,
    ) -> None:
        self.type_column: int = normalize_column(type_column)
        self.score_column: int = normalize_column(score_column)
        self.positive_value = str(positive_value).strip()
        self.negative_value = str(negative_value).strip()
        self.auto_correction_value = auto_correction_value
        self.valid_scores: frozenset[int] = frozenset(int(s) for s in valid_scores)
        self.header_row = header_row
        self.row_anchor_column: int | None = (
            normalize_column(row_anchor_column) if row_anchor_column else None
        )
        self.scan_min_col = normalize_column(scan_min_col)
        self.scan_max_col = normalize_column(scan_max_col)
        self.ranges: dict[str, SheetRanges] = ranges or {}

    def run(self, ctx: CheckContext) -> list[Difference]:
        diffs: list[Difference] = []
        score_letter = get_column_letter(self.score_column)
        for sheet_name, sr in self.ranges.items():
            if sr.skip or sheet_name not in ctx.workbook.sheetnames:
                continue
            # Both columns must be in range for the row to be meaningful.
            if not (sr.col_in_range(self.type_column) and sr.col_in_range(self.score_column)):
                continue
            ws = ctx.workbook[sheet_name]
            last_row = find_last_data_row(
                ws,
                start_row=self.header_row + 1,
                scan_min_col=self.scan_min_col,
                scan_max_col=self.scan_max_col,
            )
            lo, hi = sr.clamp_rows(self.header_row + 1, last_row)
            for row in range(lo, hi + 1):
                if not sr.row_in_range(row):
                    continue
                if self.row_anchor_column is not None and is_empty(
                    ws.cell(row=row, column=self.row_anchor_column).value
                ):
                    continue
                type_value = ws.cell(row=row, column=self.type_column).value
                if is_empty(type_value):
                    continue
                type_value = str(type_value).strip()
                score_value = ws.cell(row=row, column=self.score_column).value
                score_coord = f"{score_letter}{row}"

                if type_value.casefold() == self.positive_value.casefold():
                    if score_value != self.auto_correction_value:
                        diffs.append(
                            Difference(
                                category=DiffCategory.INCONSISTENT_TYPE_SCORE,
                                sheet=sheet_name,
                                cell=score_coord,
                                attribute=type_value,
                                golden=self.auto_correction_value,
                                other=score_value,
                                correction=self.auto_correction_value,
                                check_name=self.name
                            )
                        )
                elif type_value.casefold() == self.negative_value.casefold():
                    if not _is_valid_score(score_value, self.valid_scores):
                        diffs.append(
                            Difference(
                                category=DiffCategory.INCONSISTENT_TYPE_SCORE,
                                sheet=sheet_name,
                                cell=score_coord,
                                attribute=type_value,
                                golden=sorted(self.valid_scores),
                                other=score_value,
                                check_name=self.name,
                            )
                        )
        return diffs


def _is_valid_score(value: object, valid_scores: frozenset[int]) -> bool:
    """True iff *value* is a whole-number integer in *valid_scores*."""
    if value is None or isinstance(value, bool):
        return False
    if isinstance(value, int):
        return value in valid_scores
    if isinstance(value, float):
        return value.is_integer() and int(value) in valid_scores
    return False
