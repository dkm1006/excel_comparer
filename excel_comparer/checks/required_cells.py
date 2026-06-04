"""Required-cells check.

For every cell inside the configured ``ranges``, asserts that the cell is
non-empty. Emits a :data:`DiffCategory.MISSING_REQUIRED` finding per
empty cell.

The check runs on every sheet listed in ``ranges``. If ``ranges`` is
empty or missing the check is a no-op.

Configuration (TOML)
--------------------

.. code-block:: toml

    [[checks]]
    name = "required_cells"

    [checks.ranges]
    Impacts = ["H3:P", "U3:U"]
"""

from typing import ClassVar

from openpyxl.utils import get_column_letter

from excel_comparer.checks.base import CheckContext, register_check
from excel_comparer.checks.helpers import find_last_data_row, is_empty, normalize_column
from excel_comparer.config import SheetRanges
from excel_comparer.models import DiffCategory, Difference


@register_check
class RequiredCellsCheck:
    """Assert that every cell inside the configured ranges is non-empty."""

    name: ClassVar[str] = "required_cells"
    description: ClassVar[str] = (
        "Asserts that every cell inside the configured ranges has been filled in."
    )

    def __init__(
        self,
        *,
        header_row: int = 2,
        row_anchor_column: int | str | None = None,
        ranges: dict[str, SheetRanges] | None = None,
    ) -> None:
        self.header_row = header_row
        self.row_anchor_column: int | None = (
            normalize_column(row_anchor_column) if row_anchor_column else None
        )
        self.ranges: dict[str, SheetRanges] = ranges or {}

    def run(self, ctx: CheckContext) -> list[Difference]:
        diffs: list[Difference] = []
        for sheet_name, sr in self.ranges.items():
            if sr.skip or sheet_name not in ctx.workbook.sheetnames:
                continue
            ws = ctx.workbook[sheet_name]

            scan_min_col, scan_max_col = sr.clamp_cols(1, ws.max_column)
            if scan_min_col > scan_max_col:
                continue
            last_row = find_last_data_row(
                ws,
                start_row=self.header_row + 1,
                scan_min_col=scan_min_col,
                scan_max_col=scan_max_col,
            )
            if last_row < self.header_row + 1:
                continue

            # Per-row cache so we don't re-read the anchor cell once per
            # column hit.
            anchor_empty: dict[int, bool] = {}

            for row, col in sr.iter_cells(
                row_floor=self.header_row + 1,
                row_ceiling=last_row,
                col_ceiling=ws.max_column,
            ):
                if self.row_anchor_column is not None:
                    if row not in anchor_empty:
                        anchor_empty[row] = is_empty(
                            ws.cell(row=row, column=self.row_anchor_column).value
                        )
                    if anchor_empty[row]:
                        continue
                value = ws.cell(row=row, column=col).value
                if is_empty(value):
                    diffs.append(
                        Difference(
                            category=DiffCategory.MISSING_REQUIRED,
                            sheet=sheet_name,
                            cell=f"{get_column_letter(col)}{row}",
                            check_name=self.name,
                        )
                    )
        return diffs
